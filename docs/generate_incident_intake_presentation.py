from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUTPUT = Path(__file__).with_name("Incident_Intake_Module_Presentation.pptx")

GREEN = "174C42"
GREEN_DARK = "103B34"
GREEN_PALE = "E6F0EB"
PAPER = "F4F5EF"
WHITE = "FFFEFA"
INK = "182522"
MUTED = "78837D"
LINE = "E1E5DC"
AMBER = "E9AD43"
TERRA = "DF7B41"
RED = "A64135"
RED_PALE = "F9E9E5"


def color(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, radius=False, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = color(line)
        shape.line.width = Pt(0.8)
    if radius:
        shape.adjustments[0] = 0.08
    return shape


def text(slide, value, x, y, w, h, size=14, fill=INK, bold=False,
         font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
         margin=0, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color(fill)
    return box


def line(slide, x1, y1, x2, y2, fill=LINE, width=1.0):
    shape = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    shape.line.color.rgb = color(fill)
    shape.line.width = Pt(width)
    return shape


def base_slide(prs, number, section, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(PAPER)
    text(slide, "INCIDENT INTAKE  /  REPORT & MEDIA", 0.68, 0.34, 6.2, 0.24,
         9, GREEN, True)
    text(slide, section.upper(), 10.1, 0.34, 2.55, 0.24, 9, MUTED, True,
         align=PP_ALIGN.RIGHT)
    line(slide, 0.68, 0.72, 12.65, 0.72, LINE, 0.8)
    text(slide, title, 0.68, 0.98, 11.95, 0.58, 28, INK, True, "Aptos Display")
    text(slide, subtitle, 0.7, 1.65, 11.9, 0.42, 12, MUTED)
    line(slide, 0.68, 7.08, 12.65, 7.08, LINE, 0.8)
    text(slide, "CURRENT SCOPE  ·  EPIC 1 / STORIES 1.1 + 1.2", 0.68, 7.17,
         8.5, 0.18, 8, MUTED, True)
    text(slide, f"{number:02d}  /  07", 11.2, 7.15, 1.45, 0.2, 9, GREEN,
         True, align=PP_ALIGN.RIGHT)
    return slide


def label(slide, value, x, y, w=2.0, fill=GREEN):
    text(slide, value.upper(), x, y, w, 0.22, 9, fill, True)


def bullet(slide, value, x, y, w, size=13, bullet_color=TERRA, text_color=INK):
    rect(slide, x, y + 0.12, 0.07, 0.07, bullet_color, radius=True)
    text(slide, value, x + 0.2, y, w - 0.2, 0.48, size, text_color)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Incident Intake Module"
    prs.core_properties.subject = "Epic 1 user stories 1.1 and 1.2"
    prs.core_properties.author = "Incident Response Project Team"
    prs.core_properties.keywords = "incident reporting, JSON, CSV, media upload"

    # 1. Cover
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color(GREEN_DARK)
    rect(slide, 0, 0, 0.17, 7.5, AMBER)
    rect(slide, 9.82, 0, 3.513, 7.5, GREEN)
    text(slide, "PROJECT WALKTHROUGH  /  EPIC 1", 0.88, 0.72, 5.9, 0.28,
         10, "D8E5DD", True)
    text(slide, "INCIDENT\nINTAKE", 0.84, 1.58, 8.35, 1.85, 43, WHITE,
         True, "Aptos Display")
    rect(slide, 0.9, 3.78, 0.52, 0.06, AMBER)
    text(slide, "Report & Media Module", 0.88, 4.04, 7.8, 0.52, 23,
         "F3C96F", False, "Aptos Display")
    text(slide, "Web submission  +  structured JSON / CSV import", 0.9, 4.78,
         7.9, 0.42, 15, "D8E5DD")
    text(slide, "User stories 1.1 and 1.2  ·  Current sprint scope", 0.9, 6.77,
         7.7, 0.28, 10, "B7C9C0", True)
    text(slide, "01", 10.42, 1.24, 1.5, 0.5, 15, "F3C96F", True)
    text(slide, "TWO\nENTRY\nPOINTS", 10.38, 1.92, 2.5, 1.8, 25,
         WHITE, True, "Aptos Display")
    line(slide, 10.42, 4.32, 12.2, 4.32, "6B9487", 1.1)
    text(slide, "One report contract.\nIndependent record validation.",
         10.42, 4.56, 2.45, 0.75, 12, "D8E5DD")

    # 2. Requirement scope
    slide = base_slide(prs, 2, "What this delivers", "Two supported ways to report",
                       "Both paths feed the same report and media services in the backend.")
    cards = [
        (0.72, "01", "WEB FORM", "Enter one incident at a time.",
         ["Required details: type, description, coordinates", "Location name and image are optional"]),
        (6.82, "02", "JSON / CSV", "Submit a file containing multiple reports.",
         ["JSON objects or CSV rows use the report fields", "Each record receives its own result"]),
    ]
    for x, number, title, description, points in cards:
        rect(slide, x, 2.34, 5.78, 3.25, WHITE, radius=True, line=LINE)
        text(slide, number, x + 0.25, 2.63, 0.6, 0.38, 14, TERRA, True)
        label(slide, title, x + 0.95, 2.69, 3.4)
        text(slide, description, x + 0.27, 3.22, 5.1, 0.55, 18, INK, True,
             "Aptos Display")
        bullet(slide, points[0], x + 0.28, 4.04, 5.15, 12)
        bullet(slide, points[1], x + 0.28, 4.66, 5.15, 12)
    rect(slide, 0.72, 5.92, 11.88, 0.68, GREEN_PALE, radius=True)
    text(slide, "ON SUCCESS", 0.98, 6.15, 1.2, 0.18, 9, GREEN, True)
    text(slide, "A stored report gets a unique ID, timestamp, and SUBMITTED status.",
         2.25, 6.08, 9.8, 0.34, 13, GREEN, True)

    # 3. Web form
    slide = base_slide(prs, 3, "Story 1.1", "Web form: one report, clearly validated",
                       "The form collects structured incident details and sends them to POST /reports.")
    rect(slide, 0.72, 2.32, 7.45, 4.18, WHITE, radius=True, line=LINE)
    label(slide, "Required", 1.02, 2.64)
    fields = [
        ("INCIDENT TYPE", "FLOOD / FIRE / EARTHQUAKE / ..."),
        ("DESCRIPTION", "What is happening? Who is affected?"),
        ("LATITUDE", "-90 to 90"),
        ("LONGITUDE", "-180 to 180"),
    ]
    for index, (name, hint) in enumerate(fields):
        y = 3.06 + index * 0.7
        text(slide, name, 1.04, y, 1.56, 0.2, 8, MUTED, True)
        rect(slide, 2.55, y - 0.08, 5.18, 0.43, PAPER, radius=True)
        text(slide, hint, 2.72, y + 0.025, 4.8, 0.22, 10, MUTED)
    line(slide, 1.03, 5.94, 7.78, 5.94, LINE, 0.8)
    text(slide, "OPTIONAL", 1.04, 6.1, 0.85, 0.18, 8, TERRA, True)
    text(slide, "Location name  ·  Photo evidence", 2.06, 6.06, 4.5, 0.24,
         11, INK, True)

    rect(slide, 8.52, 2.32, 4.08, 4.18, GREEN, radius=True)
    label(slide, "Validation", 8.86, 2.66, fill="F3C96F")
    bullet(slide, "Required values and coordinate ranges are checked.",
           8.88, 3.16, 3.35, 12, AMBER, WHITE)
    bullet(slide, "Description is limited to 5,000 characters.",
           8.88, 4.12, 3.35, 12, AMBER, WHITE)
    bullet(slide, "The API rejects invalid input without saving a report.",
           8.88, 5.08, 3.35, 12, AMBER, WHITE)

    # 4. Optional media
    slide = base_slide(prs, 4, "Story 1.2", "Photo evidence is optional",
                       "Text-only reports remain valid; a supported image is safely attached to its report.")
    rect(slide, 0.72, 2.34, 5.62, 3.78, GREEN, radius=True)
    label(slide, "Accepted formats", 1.05, 2.72, fill="F3C96F")
    for idx, (kind, short) in enumerate([("JPEG", "JPG"), ("PNG", "PNG"), ("WEBP", "WEBP")]):
        x = 1.05 + idx * 1.62
        rect(slide, x, 3.28, 1.35, 0.87, "245B4F", radius=True)
        text(slide, short, x, 3.54, 1.35, 0.32, 18, WHITE, True,
             align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        text(slide, kind, x, 4.34, 1.35, 0.22, 9, "D8E5DD", True,
             align=PP_ALIGN.CENTER)
    text(slide, "MAXIMUM SIZE", 1.05, 5.14, 1.35, 0.2, 8, "B7C9C0", True)
    text(slide, "5 MB", 2.45, 5.04, 1.3, 0.45, 24, "F3C96F", True,
         "Aptos Display")
    text(slide, "Checked by the backend", 1.05, 5.64, 4.5, 0.25,
         11, "D8E5DD")

    rect(slide, 6.7, 2.34, 5.9, 3.78, WHITE, radius=True, line=LINE)
    label(slide, "Storage flow", 7.05, 2.72)
    flow = [("SELECT", "Optional file"), ("VALIDATE", "Type + bytes"),
            ("STORE", "Safe filename"), ("LINK", "Report media")]
    for idx, (heading, detail) in enumerate(flow):
        y = 3.22 + idx * 0.61
        rect(slide, 7.05, y, 1.14, 0.35, GREEN_PALE, radius=True)
        text(slide, heading, 7.05, y + 0.08, 1.14, 0.17, 8, GREEN, True,
             align=PP_ALIGN.CENTER)
        text(slide, detail, 8.48, y + 0.03, 3.6, 0.24, 12, INK, True)
        if idx < len(flow) - 1:
            line(slide, 7.62, y + 0.37, 7.62, y + 0.58, "9DB5A8", 1.0)
    rect(slide, 0.72, 6.32, 11.88, 0.42, "FBF1D9", radius=True)
    text(slide, "No image selected? The valid report still submits normally.",
         0.96, 6.42, 11.35, 0.19, 11, INK, True)

    # 5. Structured import
    slide = base_slide(prs, 5, "Story 1.1", "Bulk import: partial success is intentional",
                       "JSON and CSV records are validated and saved independently, not as one all-or-nothing batch.")
    rect(slide, 0.72, 2.34, 4.02, 3.98, WHITE, radius=True, line=LINE)
    label(slide, "Expected fields", 1.02, 2.66)
    text(slide, "disaster_type\ndescription\nlatitude\nlongitude",
         1.04, 3.1, 2.6, 1.5, 15, GREEN, True, "Cascadia Code")
    line(slide, 1.02, 4.91, 4.35, 4.91, LINE, 0.8)
    text(slide, "Optional: location_name", 1.04, 5.14, 3.2, 0.23,
         10, MUTED, True)
    text(slide, "Up to 1,000 records by default", 1.04, 5.68, 3.3, 0.23,
         10, MUTED)

    rect(slide, 5.02, 2.34, 7.58, 3.98, WHITE, radius=True, line=LINE)
    label(slide, "Example response", 5.34, 2.66)
    rect(slide, 5.34, 3.06, 6.95, 0.42, GREEN, radius=True)
    text(slide, "ROW", 5.58, 3.18, 0.7, 0.16, 8, WHITE, True)
    text(slide, "RESULT", 6.66, 3.18, 1.1, 0.16, 8, WHITE, True)
    text(slide, "DETAIL", 8.12, 3.18, 3.2, 0.16, 8, WHITE, True)
    examples = [("01", "ACCEPTED", "Saved as its own report", GREEN_PALE, GREEN),
                ("02", "REJECTED", "Description is required", RED_PALE, RED),
                ("03", "ACCEPTED", "Saved as a separate report", GREEN_PALE, GREEN)]
    for idx, (row, status, detail, bg, fg) in enumerate(examples):
        y = 3.64 + idx * 0.62
        if idx % 2 == 0:
            rect(slide, 5.34, y - 0.04, 6.95, 0.53, PAPER)
        text(slide, row, 5.58, y + 0.08, 0.6, 0.19, 10, MUTED, True)
        rect(slide, 6.66, y + 0.02, 1.04, 0.29, bg, radius=True)
        text(slide, status, 6.66, y + 0.08, 1.04, 0.14, 7, fg, True,
             align=PP_ALIGN.CENTER)
        text(slide, detail, 8.12, y + 0.07, 3.7, 0.2, 10, INK)
    rect(slide, 5.34, 5.68, 6.95, 0.42, GREEN_PALE, radius=True)
    text(slide, "2 accepted  ·  1 rejected  ·  accepted rows persist",
         5.56, 5.79, 6.5, 0.18, 10, GREEN, True)

    # 6. Architecture and boundaries
    slide = base_slide(prs, 6, "Implementation", "Small frontend, reusable backend",
                       "The interface uses existing endpoints; domain rules and persistence stay in the backend.")
    steps = [
        ("REACT UI", "Web form\nJSON / CSV"),
        ("FASTAPI", "POST /reports\nPOST /reports/import/*"),
        ("VALIDATION", "Pydantic\nimage checks"),
        ("SERVICES", "report_service\nimport_service"),
        ("STORAGE", "SQLAlchemy\nSQLite in dev"),
    ]
    x_values = [0.72, 3.15, 5.58, 8.01, 10.44]
    for idx, ((heading, detail), x) in enumerate(zip(steps, x_values)):
        rect(slide, x, 2.72, 2.12, 1.28, WHITE, radius=True, line=LINE)
        text(slide, heading, x + 0.14, 2.96, 1.84, 0.22, 9, GREEN, True,
             align=PP_ALIGN.CENTER)
        text(slide, detail, x + 0.12, 3.33, 1.88, 0.52, 10, MUTED,
             align=PP_ALIGN.CENTER)
        if idx < len(steps) - 1:
            line(slide, x + 2.14, 3.36, x + 2.38, 3.36, TERRA, 1.7)
    rect(slide, 0.72, 4.48, 5.77, 1.47, GREEN_PALE, radius=True)
    label(slide, "This delivery", 1.0, 4.76)
    bullet(slide, "Web submission, optional image, structured imports", 1.02,
           5.13, 5.1, 11, GREEN)
    rect(slide, 6.76, 4.48, 5.84, 1.47, "FBF1D9", radius=True)
    label(slide, "Future team scope", 7.04, 4.76, fill="8A641C")
    bullet(slide, "Auth, incident review, maps, resources, AI planning", 7.06,
           5.13, 5.1, 11, TERRA)
    text(slide, "A report is SUBMITTED evidence, not a confirmed incident or dispatch.",
         0.82, 6.35, 11.65, 0.25, 12, INK, True, align=PP_ALIGN.CENTER)

    # 7. Demo and handoff
    slide = base_slide(prs, 7, "Demo", "A short walkthrough",
                       "Show the behavior that maps directly to the two completed acceptance criteria.")
    demo = [
        ("01", "Open the web form", "Point out required fields and optional location/photo."),
        ("02", "Submit without a photo", "Show a successful response and report reference ID."),
        ("03", "Choose JSON / CSV", "Show the field template and import controls."),
        ("04", "Import mixed-validity rows", "Show accepted and rejected results for each record."),
    ]
    for idx, (number, heading, detail) in enumerate(demo):
        y = 2.44 + idx * 0.82
        rect(slide, 0.78, y, 0.52, 0.52, GREEN, radius=True)
        text(slide, number, 0.78, y + 0.15, 0.52, 0.19, 10, WHITE, True,
             align=PP_ALIGN.CENTER)
        text(slide, heading, 1.55, y + 0.01, 3.2, 0.28, 14, INK, True)
        text(slide, detail, 4.62, y + 0.03, 7.65, 0.32, 11, MUTED)
        if idx < len(demo) - 1:
            line(slide, 1.04, y + 0.55, 1.04, y + 0.79, "B7C9C0", 1.1)
    rect(slide, 0.78, 5.95, 11.82, 0.68, GREEN_DARK, radius=True)
    text(slide, "TAKEAWAY", 1.05, 6.18, 1.0, 0.18, 8, "F3C96F", True)
    text(slide, "Two intake paths. One validated report pipeline. More modules remain for the team.",
         2.14, 6.12, 10.0, 0.3, 12, WHITE, True)

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()